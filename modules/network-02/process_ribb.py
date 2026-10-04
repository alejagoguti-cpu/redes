#!/usr/bin/env python3
"""Procesa RIBB para network-02 sin alterar la base original."""
from __future__ import annotations
import argparse,csv,json,re,unicodedata
from collections import Counter,defaultdict
from pathlib import Path
from openpyxl import load_workbook

SOURCE="https://redbiotica.jbb.gov.co/"
DATASET_DATE="2024-12-30"

def clean(v):
    return "" if v is None else str(v).strip()

def norm(v):
    s=clean(v).lower()
    s="".join(c for c in unicodedata.normalize("NFD",s) if unicodedata.category(c)!="Mn")
    return re.sub(r"\s+"," ",re.sub(r"[^a-z0-9×.-]+"," ",s)).strip()

def write_csv(path,rows,fields=None):
    path.parent.mkdir(parents=True,exist_ok=True)
    if not rows:
        path.write_text("",encoding="utf-8"); return
    fields=fields or list(rows[0].keys())
    with path.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction="ignore"); w.writeheader(); w.writerows(rows)

def probable_species(name):
    p=clean(name).split()
    return len(p)>=2 and p[0][:1].isupper()

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("input",nargs="?",default="assets/data/RIBB_v2024-12-30.xlsx")
    ap.add_argument("--output-dir",default="assets/data")
    args=ap.parse_args()
    src=Path(args.input); out=Path(args.output_dir)
    if not src.exists(): raise SystemExit(f"No se encontró: {src}")

    wb=load_workbook(src,read_only=True,data_only=True)
    ws=wb[wb.sheetnames[0]]
    it=ws.iter_rows(values_only=True)
    headers=[clean(v) or f"col_{i+1}" for i,v in enumerate(next(it))]
    records=[]
    taxa=defaultdict(lambda:{
        "rows":0,"roles":Counter(),"origins":Counter(),"types":Counter(),
        "titles":Counter(),"refs":set(),"media":set(),"coords":set(),"dates":set()
    })

    for rn,vals in enumerate(it,start=2):
        if not any(clean(v) for v in vals): continue
        row={h:clean(vals[i]) if i<len(vals) else "" for i,h in enumerate(headers)}
        refs=[]
        for n in range(1,10):
            ref=row.get(f"Referencia {n}","")
            typ=row.get(f"Tipo de referencia {n}","")
            med=row.get(f"Multimedia {n}","")
            coord=row.get(f"Coodenadas {n}","") or row.get(f"Coordenadas {n}","")
            fecha=row.get(f"Fecha {n}","")
            if any([ref,typ,med,coord,fecha]):
                refs.append({"n":n,"tipo":typ,"referencia":ref,"multimedia":med,"coordenadas":coord,"fecha":fecha})

        t1=row.get("Taxón 1",""); t2=row.get("Taxón 2","")
        nref=sum(bool(r["referencia"]) for r in refs)
        nmedia=sum(bool(r["multimedia"]) for r in refs)
        nloc=sum(bool(r["coordenadas"]) and norm(r["coordenadas"])!="n a" for r in refs)
        ndate=sum(bool(r["fecha"]) for r in refs)
        score=sum([bool(t1),bool(t2),bool(row.get("Tipo de interacción","")),nref>0,nloc>0,ndate>0])
        quality="Alta" if score>=5 and nref else ("Media" if score>=4 and nref else "Baja")

        rec={
            "ribb_id":f"RIBB-{rn-1:05d}","taxon_1":t1,"origen_taxon_1":row.get("Origen taxón 1",""),
            "tipo_interaccion":row.get("Tipo de interacción",""),"parte_planta":row.get("parte de la planta afectada",""),
            "taxon_2":t2,"origen_taxon_2":row.get("Origen taxón 2",""),
            "titulo_interaccion":row.get("Título de interacción",""),"reportada_distrito":row.get("Reportada en el distrito",""),
            "frecuencia":row.get("Frecuencia",""),"observaciones":row.get("Observaciones",""),
            "numero_referencias":nref,"numero_multimedia":nmedia,"numero_localizaciones":nloc,"numero_fechas":ndate,
            "puntaje_completitud_6":score,"calidad_documental":quality,
            "referencias_json":json.dumps(refs,ensure_ascii=False),
            "fuente_base":"Red de Interacciones Bióticas de Bogotá D.C. (RIBB) · JBB",
            "url_fuente_base":SOURCE,"fecha_base":DATASET_DATE,
        }
        records.append(rec)

        for role,taxon,origin in [(1,t1,row.get("Origen taxón 1","")),(2,t2,row.get("Origen taxón 2",""))]:
            if not taxon: continue
            d=taxa[taxon]; d["rows"]+=1; d["roles"][f"taxon_{role}"]+=1
            if origin:d["origins"][origin]+=1
            if rec["tipo_interaccion"]:d["types"][rec["tipo_interaccion"]]+=1
            if rec["titulo_interaccion"]:d["titles"][rec["titulo_interaccion"]]+=1
            for r in refs:
                if r["referencia"]:d["refs"].add(r["referencia"])
                if r["multimedia"]:d["media"].add(r["multimedia"])
                if r["coordenadas"] and norm(r["coordenadas"])!="n a":d["coords"].add(r["coordenadas"])
                if r["fecha"]:d["dates"].add(r["fecha"])

    wb.close()
    species=[]
    for taxon,d in sorted(taxa.items(),key=lambda kv:(-kv[1]["rows"],kv[0])):
        species.append({
            "taxon":taxon,"especie_probable":"Sí" if probable_species(taxon) else "No","registros":d["rows"],
            "como_taxon_1":d["roles"]["taxon_1"],"como_taxon_2":d["roles"]["taxon_2"],
            "origenes":" | ".join(f"{k}: {v}" for k,v in d["origins"].most_common()),
            "tipos_interaccion":" | ".join(f"{k}: {v}" for k,v in d["types"].most_common()),
            "titulos_interaccion":" | ".join(f"{k}: {v}" for k,v in d["titles"].most_common()),
            "multimedia_ribb":" | ".join(sorted(d["media"])),"n_multimedia_ribb":len(d["media"]),
            "n_referencias":len(d["refs"]),"n_coordenadas":len(d["coords"]),"n_fechas":len(d["dates"]),
            "imagen_url":"","imagen_fuente":"","imagen_estado":"Pendiente de resolución por nombre científico"
        })

    out.mkdir(parents=True,exist_ok=True)
    write_csv(out/"ribb_interacciones.csv",records)
    qfields=["ribb_id","taxon_1","taxon_2","tipo_interaccion","numero_referencias","numero_multimedia","numero_localizaciones","numero_fechas","puntaje_completitud_6","calidad_documental","fuente_base","url_fuente_base","fecha_base"]
    write_csv(out/"ribb_fuentes_calidad.csv",records,qfields)
    write_csv(out/"ribb_especies.csv",species)
    (out/"ribb_especies.json").write_text(json.dumps([r for r in species if r["especie_probable"]=="Sí"],ensure_ascii=False),encoding="utf-8")

    summary={
        "dataset_date":DATASET_DATE,"source":SOURCE,"total_interacciones":len(records),
        "taxones_unicos":len(species),"especies_probables":sum(r["especie_probable"]=="Sí" for r in species),
        "taxones_con_multimedia_ribb":sum(r["n_multimedia_ribb"]>0 for r in species),
        "tipos_interaccion":dict(Counter(r["tipo_interaccion"] for r in records)),
        "titulos_interaccion":dict(Counter(r["titulo_interaccion"] for r in records)),
        "calidad_documental":dict(Counter(r["calidad_documental"] for r in records)),
    }
    (out/"ribb_resumen.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=="__main__": main()
