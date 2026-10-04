#!/usr/bin/env python3
"""Construye dataset web RIBB con semántica y layout 3D derivados de la red real."""
import json, sys
from pathlib import Path
from collections import Counter, defaultdict
from openpyxl import load_workbook
import networkx as nx

src=Path(sys.argv[1] if len(sys.argv)>1 else "assets/data/RIBB_v2024-12-30.xlsx")
out=Path(sys.argv[2] if len(sys.argv)>2 else "assets/data/ribb_web.json")

wb=load_workbook(src,read_only=True,data_only=True)
ws=wb[wb.sheetnames[0]]
it=ws.iter_rows(values_only=True)
headers=[str(v).strip() if v is not None else "" for v in next(it)]
idx={h:i for i,h in enumerate(headers)}

def val(row,key):
    i=idx.get(key)
    return "" if i is None or i>=len(row) or row[i] is None else str(row[i]).strip()

def refs(row):
    out=[]
    for n in range(1,10):
        ref=val(row,f"Referencia {n}")
        typ=val(row,f"Tipo de referencia {n}")
        med=val(row,f"Multimedia {n}")
        coord=val(row,f"Coodenadas {n}") or val(row,f"Coordenadas {n}")
        fecha=val(row,f"Fecha {n}")
        if any((ref,typ,med,coord,fecha)):
            out.append({"tipo":typ,"ref":ref,"media":med,"coord":coord,"fecha":fecha})
    return out

DIRECTED={
    "Consume","Crece sobre","Oviposita en","Parasita a","Dispersa",
    "Dispersa a","Poliniza a"
}
SYMMETRIC={
    "Es simbionte con","Se agrede con","Alelopatía negativa con",
    "Alelopatía positiva con"
}

def semantic_phrase(verb,actor,target,part):
    p=(part or "").strip()
    if verb=="Consume":
        return f"{actor} consume {p + ' de ' if p and p.lower()!='general' else ''}{target}."
    if verb=="Crece sobre": return f"{actor} crece sobre {target}."
    if verb=="Oviposita en": return f"{actor} oviposita en {target}."
    if verb=="Parasita a": return f"{actor} parasita a {target}."
    if verb=="Dispersa": return f"{actor} dispersa estructuras reproductivas de {target}."
    if verb=="Dispersa a": return f"{actor} dispersa a {target}."
    if verb=="Poliniza a": return f"{actor} poliniza a {target}."
    if verb=="Es simbionte con": return f"{actor} mantiene una relación simbiótica con {target}."
    if verb=="Nidifica con": return f"{actor} nidifica con o utiliza recursos asociados a {target}."
    if verb=="Se agrede con": return f"{actor} presenta una interacción agonística con {target}."
    if verb.startswith("Alelopatía"): return f"{actor} presenta {verb.lower()} {target}."
    return f"{actor} interactúa con {target}."

records=[]
node_meta=defaultdict(lambda:{"n":0,"origins":Counter(),"categories":Counter()})
for rn,row in enumerate(it,start=2):
    t1=val(row,"Taxón 1"); t2=val(row,"Taxón 2")
    if not t1 or not t2: continue
    verb=val(row,"Tipo de interacción")
    category=val(row,"Título de interacción") or "OTRAS INTERACCIONES"
    rr=refs(row)
    first=rr[0] if rr else {"tipo":"","ref":"","media":"","coord":"","fecha":""}
    district="SI" if val(row,"Reportada en el distrito").upper()=="SI" else "NO"
    direction="directed" if verb in DIRECTED else ("symmetric" if verb in SYMMETRIC else "contextual")
    rec={
        "id":f"R{rn-1}",
        "source":t1,"target":t2,
        "verb":verb,"category":category,
        "direction":direction,
        "phrase":semantic_phrase(verb,t1,t2,val(row,"parte de la planta afectada")),
        "part":val(row,"parte de la planta afectada"),
        "sourceOrigin":val(row,"Origen taxón 1"),"targetOrigin":val(row,"Origen taxón 2"),
        "district":district,
        "frequency":val(row,"Frecuencia"),"observations":val(row,"Observaciones"),
        "reference":first["ref"],"referenceType":first["tipo"],"media":first["media"],
        "coord":first["coord"],"date":first["fecha"],
        "referenceCount":sum(bool(x["ref"]) for x in rr)
    }
    score=sum(bool(rec[k]) for k in ["source","target","verb","reference","coord","date"])
    rec["quality"]="Alta" if rec["reference"] and score>=5 else ("Media" if rec["reference"] and score>=4 else "Baja")
    records.append(rec)
    for taxon,origin in ((t1,rec["sourceOrigin"]),(t2,rec["targetOrigin"])):
        node_meta[taxon]["n"]+=1
        if origin: node_meta[taxon]["origins"][origin]+=1
        node_meta[taxon]["categories"][category]+=1

def make_graph(rows):
    G=nx.Graph()
    for r in rows:
        a,b=r["source"],r["target"]
        if G.has_edge(a,b): G[a][b]["weight"]+=1
        else: G.add_edge(a,b,weight=1)
    return G

def layout3d(G,seed):
    if not G.nodes: return {}
    pos=nx.spring_layout(G,dim=3,seed=seed,iterations=80,weight="weight",scale=1.0)
    return {k:[round(float(x),6) for x in v] for k,v in pos.items()}

district_rows=[r for r in records if r["district"]=="SI"]
Gd=make_graph(district_rows)
Ga=make_graph(records)
pos_d=layout3d(Gd,42)
pos_a=layout3d(Ga,43)

nodes=[]
for name,m in node_meta.items():
    origin=m["origins"].most_common(1)[0][0] if m["origins"] else ""
    dominant=m["categories"].most_common(1)[0][0] if m["categories"] else "OTRAS INTERACCIONES"
    nodes.append({
        "id":name,"n":m["n"],"origin":origin,"dominantCategory":dominant,
        "positionDistrict":pos_d.get(name),
        "positionAll":pos_a.get(name)
    })

payload={
    "meta":{
        "source":"https://redbiotica.jbb.gov.co/",
        "dataset":"RIBB_v2024-12-30.xlsx",
        "totalRecords":len(records),
        "districtRecords":len(district_rows),
        "taxa":len(nodes),
        "districtTaxa":Gd.number_of_nodes(),
        "districtComponents":nx.number_connected_components(Gd),
        "interactionCategories":dict(Counter(r["category"] for r in records)),
        "interactionVerbs":dict(Counter(r["verb"] for r in records))
    },
    "nodes":nodes,
    "records":records
}
out.parent.mkdir(parents=True,exist_ok=True)
out.write_text(json.dumps(payload,ensure_ascii=False,separators=(",",":")),encoding="utf-8")
print(json.dumps(payload["meta"],ensure_ascii=False,indent=2))
